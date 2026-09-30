from io import BytesIO
from typing import Protocol

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.documents.schemas import DocumentKind

PDF_HEADER_PREFIX_LIMIT = 32
_PDF_BOM = b"\xef\xbb\xbf"
_PDF_ALLOWED_PREFIX_BYTES = b" \t\r\n\f"


def prepare_pdf_parser_input(content: bytes) -> tuple[bytes, int]:
    """Return a bounded derived view for PdfReader without mutating raw bytes."""

    prefix = content[:PDF_HEADER_PREFIX_LIMIT]
    offset = 0
    if prefix.startswith(_PDF_BOM):
        offset = len(_PDF_BOM)
    while offset < len(prefix) and prefix[offset] in _PDF_ALLOWED_PREFIX_BYTES:
        offset += 1
    if content[offset : offset + 5] != b"%PDF-":
        raise ValueError("invalid_pdf_header")
    return content[offset:], offset


class ParsedDocument(Protocol):
    document_id: str
    kind: DocumentKind
    content: str


def parse_document(content: bytes, content_type: str) -> str:
    if content_type == "application/pdf":
        text = "\n".join(parse_pdf_pages(content))
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        text = "\n".join(paragraph.text for paragraph in DocxDocument(BytesIO(content)).paragraphs)
    else:
        raise ValueError("unsupported_document_format")
    normalized = text.strip()
    if not normalized:
        raise ValueError("document_text_empty")
    return normalized


def parse_pdf_pages(content: bytes) -> list[str]:
    parser_input, _ = prepare_pdf_parser_input(content)
    return [page.extract_text() or "" for page in PdfReader(BytesIO(parser_input)).pages]
