# Copyright (c) 2021, FOSS United and contributors
# For license information, please see license.txt

"""Import / Export of quizzes as XLSX.

The spreadsheet is written for the person building the quiz, not for the database.
Sheet "Quiz Questions" holds one row per question with readable headers; sheet
"Quiz Settings" holds one row per quiz so settings are not repeated on every
question row.

Writing always goes through `Document.save()` / `Document.insert()` so every
controller validation in LMS Quiz and LMS Question still applies.
"""

import html
import re
from io import BytesIO

import frappe
import openpyxl
from frappe import _
from frappe.utils import cint, cstr
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from lms.lms.utils import has_course_instructor_role, has_course_moderator_role

QUESTIONS_SHEET = "Quiz Questions"
SETTINGS_SHEET = "Quiz Settings"
INSTRUCTIONS_SHEET = "Instructions"

QUESTION_TYPES = ("Choices", "User Input", "Open Ended")

# Several accepted answers / several explanations live in one cell.
MULTI_SEPARATOR = "|"

# Header labels of the questions sheet, in display order.
QUESTION_COLUMNS = (
	"Quiz Title",
	"Question",
	"Question Type",
	"Marks",
	"Option 1",
	"Option 2",
	"Option 3",
	"Option 4",
	"Correct Answer",
	"Expected Answer",
	"Explanation",
)

# Header label -> internal key used while parsing.
QUESTION_KEYS = {
	"quiz title": "quiz_title",
	"question": "question",
	"question type": "question_type",
	"marks": "marks",
	"option 1": "option_1",
	"option 2": "option_2",
	"option 3": "option_3",
	"option 4": "option_4",
	"correct answer": "correct_answer",
	"expected answer": "expected_answer",
	"explanation": "explanation",
}

REQUIRED_QUESTION_COLUMNS = ("Quiz Title", "Question", "Question Type", "Marks")

# Header labels of the settings sheet -> LMS Quiz fieldname.
SETTINGS_COLUMNS = (
	("Quiz Title", "title"),
	("Passing Percentage", "passing_percentage"),
	("Maximum Attempts", "max_attempts"),
	("Duration", "duration"),
	("Show Answers", "show_answers"),
	("Show Submission History", "show_submission_history"),
	("Shuffle Questions", "shuffle_questions"),
	("Limit Questions To", "limit_questions_to"),
	("Enable Negative Marking", "enable_negative_marking"),
	("Marks To Cut", "marks_to_cut"),
)
SETTINGS_FIELDS = {label.lower(): fieldname for label, fieldname in SETTINGS_COLUMNS}

SETTINGS_CHECKBOXES = (
	"show_answers",
	"show_submission_history",
	"shuffle_questions",
	"enable_negative_marking",
)
SETTINGS_INTS = (
	"passing_percentage",
	"max_attempts",
	"limit_questions_to",
	"marks_to_cut",
)

COLUMN_WIDTHS = {
	"Quiz Title": 28,
	"Question": 50,
	"Question Type": 16,
	"Marks": 8,
	"Correct Answer": 16,
	"Expected Answer": 34,
	"Explanation": 40,
	"Passing Percentage": 20,
	"Maximum Attempts": 20,
	"Show Submission History": 24,
	"Enable Negative Marking": 24,
	"Limit Questions To": 20,
}
DEFAULT_COLUMN_WIDTH = 22

MAX_ROWS = 2000
MAX_ERRORS = 100

TRUTHY = {"1", "yes", "y", "true", "x", "có", "co"}


def check_permission():
	if not (has_course_moderator_role() or has_course_instructor_role()):
		frappe.throw(_("You are not permitted to import or export quizzes."), frappe.PermissionError)


def check_quiz_permission(quiz_name):
	"""Mirrors the Quizzes list: moderators see everything, instructors only their own."""
	if has_course_moderator_role():
		return

	if frappe.db.get_value("LMS Quiz", quiz_name, "owner") != frappe.session.user:
		frappe.throw(_("You are not permitted to export this quiz."), frappe.PermissionError)


def plain_text(value):
	"""Strips HTML from a Text Editor value so it reads well in a spreadsheet."""
	text = cstr(value)
	text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
	text = re.sub(r"</(p|div|li|h[1-6])>", " ", text, flags=re.IGNORECASE)
	text = re.sub(r"<[^>]+>", "", text)
	text = html.unescape(text)
	return re.sub(r"\s+", " ", text).strip()


def normalize(value):
	"""Key used to decide whether two questions are the same question."""
	return plain_text(value).lower()


def header_key(value):
	return re.sub(r"\s+", " ", cstr(value)).strip().lower()


def to_bool(value):
	if value is None or value == "":
		return 0

	if isinstance(value, bool):
		return 1 if value else 0

	if isinstance(value, (int, float)):
		return 1 if int(value) else 0

	return 1 if cstr(value).strip().lower() in TRUTHY else 0


def to_yes_no(value):
	return _("Yes") if cint(value) else _("No")


def cell_value(value):
	"""openpyxl gives numbers as float; keep them printable and trim strings."""
	if value is None:
		return ""

	if isinstance(value, float) and value.is_integer():
		return str(int(value))

	if isinstance(value, str):
		return value.strip()

	return cstr(value).strip()


def split_multi(value):
	return [part.strip() for part in cstr(value).split(MULTI_SEPARATOR) if part.strip()]


def write_sheet(ws, columns, rows):
	ws.append(list(columns))

	for index, column in enumerate(columns, start=1):
		ws.cell(row=1, column=index).font = Font(bold=True)
		ws.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS.get(
			column, DEFAULT_COLUMN_WIDTH
		)

	for row in rows:
		ws.append([row.get(column, "") for column in columns])

	ws.freeze_panes = "A2"


def build_workbook(question_rows, settings_rows):
	wb = openpyxl.Workbook()

	questions = wb.active
	questions.title = QUESTIONS_SHEET
	write_sheet(questions, QUESTION_COLUMNS, question_rows)

	settings = wb.create_sheet(SETTINGS_SHEET)
	write_sheet(settings, [label for label, _fieldname in SETTINGS_COLUMNS], settings_rows)

	notes = wb.create_sheet(INSTRUCTIONS_SHEET)
	notes.column_dimensions["A"].width = 110
	for index, line in enumerate(get_instructions(), start=1):
		cell = notes.cell(row=index, column=1, value=line)
		cell.alignment = Alignment(wrap_text=True, vertical="top")
		if index == 1:
			cell.font = Font(bold=True)

	stream = BytesIO()
	wb.save(stream)
	return stream.getvalue()


def get_instructions():
	return [
		_("How to fill in this file"),
		"",
		_("Sheet '{0}' — one row per question.").format(QUESTIONS_SHEET),
		_("Sheet '{0}' — one row per quiz. Optional.").format(SETTINGS_SHEET),
		_("Do not rename, reorder or delete the header rows."),
		"",
		_("Always required: Quiz Title, Question, Question Type, Marks."),
		_("Rows with the same Quiz Title are imported into the same quiz."),
		_("Question Type must be one of: {0}.").format(", ".join(QUESTION_TYPES)),
		_("Marks must be a whole number greater than 0."),
		"",
		_("Question Type = Choices"),
		_("- Fill Option 1 and Option 2. Option 3 and Option 4 are optional."),
		_("- Options within one question must not repeat."),
		_("- Correct Answer holds the option numbers that are correct, e.g. 1 or 1,3."),
		_("- Leave Expected Answer empty."),
		_("- Explanation is shown under the option it belongs to, so write it as '1: text'."),
		_("  Use '{0}' between explanations, e.g. '1: Correct {0} 3: Also correct'.").format(
			MULTI_SEPARATOR
		),
		"",
		_("Question Type = User Input"),
		_("- Expected Answer holds the accepted answers. Leave Options and Correct Answer empty."),
		_("- Separate up to 4 accepted answers with '{0}', e.g. 'CPU {0} Central Processing Unit'.").format(
			MULTI_SEPARATOR
		),
		_("- Answers are matched loosely, so small typing differences still count as correct."),
		"",
		_("Question Type = Open Ended"),
		_("- Leave Options, Correct Answer and Expected Answer empty. These are graded manually."),
		_("- Explanation can hold a grading note for the reviewer."),
		_("- A quiz cannot mix Open Ended questions with other question types."),
		"",
		_("Sheet '{0}'").format(SETTINGS_SHEET),
		_("- Quiz Title must match a title used in the questions sheet."),
		_("- Yes/No columns accept Yes, No, 1 or 0."),
		_("- Settings are applied only when the quiz is created."),
		_("- For a quiz that already exists, this sheet is ignored and its settings are kept."),
		"",
		_("Existing quizzes"),
		_("- If a quiz with the same title already exists, the questions are added to it."),
		_("- A question whose text already exists in that quiz is skipped, not duplicated."),
		"",
		_("The file is checked before anything is saved. Nothing is imported if there are errors."),
		_("Maximum {0} question rows per file.").format(MAX_ROWS),
	]


def provide_file(content, filename):
	from frappe.desk.utils import provide_binary_file

	provide_binary_file(filename, "xlsx", content)


@frappe.whitelist(methods=["GET"])
def download_template():
	"""Empty template with one worked example of each question type."""
	check_permission()

	title = _("Sample Quiz")

	questions = [
		{
			"Quiz Title": title,
			"Question": _("Which of these are programming languages?"),
			"Question Type": "Choices",
			"Marks": 1,
			"Option 1": "Python",
			"Option 2": "HTML",
			"Option 3": "JavaScript",
			"Option 4": "CSS",
			"Correct Answer": "1,3",
			"Explanation": _("1: Python is a programming language. {0} 3: So is JavaScript.").format(
				MULTI_SEPARATOR
			),
		},
		{
			"Quiz Title": title,
			"Question": _("What does CPU stand for?"),
			"Question Type": "User Input",
			"Marks": 1,
			"Expected Answer": f"Central Processing Unit {MULTI_SEPARATOR} CPU",
		},
		{
			"Quiz Title": title,
			"Question": _("Explain Object Oriented Programming."),
			"Question Type": "Open Ended",
			"Marks": 5,
			"Explanation": _("Student should mention encapsulation, inheritance and polymorphism."),
		},
	]

	settings = [
		{
			"Quiz Title": title,
			"Passing Percentage": 60,
			"Maximum Attempts": 2,
			"Duration": "",
			"Show Answers": to_yes_no(1),
			"Show Submission History": to_yes_no(0),
			"Shuffle Questions": to_yes_no(0),
			"Limit Questions To": "",
			"Enable Negative Marking": to_yes_no(0),
			"Marks To Cut": "",
		}
	]

	provide_file(build_workbook(questions, settings), "quiz_import_template")


def format_correct_answer(question):
	return ",".join(str(num) for num in range(1, 5) if question.get(f"is_correct_{num}"))


def format_expected_answer(question):
	answers = [
		plain_text(question.get(f"possibility_{num}"))
		for num in range(1, 5)
		if question.get(f"possibility_{num}")
	]
	return f" {MULTI_SEPARATOR} ".join(answers)


def format_explanation(question):
	"""explanation_N belongs to option N, so the merged cell keeps the option number."""
	parts = []
	for num in range(1, 5):
		text = plain_text(question.get(f"explanation_{num}"))
		if not text:
			continue
		parts.append(f"{num}: {text}" if question.type == "Choices" else text)

	return f" {MULTI_SEPARATOR} ".join(parts)


def get_quiz_rows(quiz_name):
	quiz = frappe.get_doc("LMS Quiz", quiz_name)

	question_names = [row.question for row in quiz.questions]
	questions = {}
	if question_names:
		fields = ["name", "question", "type"]
		for num in range(1, 5):
			fields += [f"option_{num}", f"is_correct_{num}", f"explanation_{num}", f"possibility_{num}"]

		for question in frappe.get_all("LMS Question", filters={"name": ["in", question_names]}, fields=fields):
			questions[question.name] = question

	rows = []
	for row in quiz.questions:
		question = questions.get(row.question)
		if not question:
			continue

		data = {
			"Quiz Title": quiz.title,
			"Question": plain_text(question.question),
			"Question Type": question.type,
			"Marks": cint(row.marks),
			"Explanation": format_explanation(question),
		}

		if question.type == "Choices":
			for num in range(1, 5):
				data[f"Option {num}"] = plain_text(question.get(f"option_{num}"))
			data["Correct Answer"] = format_correct_answer(question)
		elif question.type == "User Input":
			data["Expected Answer"] = format_expected_answer(question)

		rows.append(data)

	return quiz, rows, [get_settings_row(quiz)]


# 0 means "not set" for these, so an empty cell reads better than a 0.
BLANK_WHEN_ZERO = ("limit_questions_to", "marks_to_cut")


def get_settings_row(quiz):
	row = {}

	for label, fieldname in SETTINGS_COLUMNS:
		value = quiz.get(fieldname)

		if fieldname in SETTINGS_CHECKBOXES:
			row[label] = to_yes_no(value)
		elif value in (None, "") or (fieldname in BLANK_WHEN_ZERO and not cint(value)):
			row[label] = ""
		else:
			row[label] = value

	return row


@frappe.whitelist(methods=["GET"])
def export_quiz(quiz):
	check_permission()

	if not frappe.db.exists("LMS Quiz", quiz):
		frappe.throw(_("Quiz {0} not found.").format(quiz), frappe.DoesNotExistError)

	check_quiz_permission(quiz)

	quiz_doc, rows, settings = get_quiz_rows(quiz)
	filename = re.sub(r"[^A-Za-z0-9_-]+", "_", cstr(quiz_doc.title)).strip("_") or quiz_doc.name

	provide_file(build_workbook(rows, settings), filename)


def read_workbook(file_url):
	"""Returns the raw rows of the questions sheet and of the settings sheet."""
	file_doc = frappe.db.get_value("File", {"file_url": file_url}, ["name", "file_name"], as_dict=True)
	if not file_doc:
		frappe.throw(_("Uploaded file not found. Please upload the file again."))

	if not cstr(file_doc.file_name).lower().endswith(".xlsx"):
		frappe.throw(_("Only .xlsx files are supported. Please upload an Excel file."))

	path = frappe.get_doc("File", file_doc.name).get_full_path()

	try:
		wb = openpyxl.load_workbook(filename=path, read_only=True, data_only=True)
	except Exception:
		frappe.throw(_("This file could not be read as an Excel workbook. Please use the template."))

	try:
		if QUESTIONS_SHEET in wb.sheetnames:
			questions_ws = wb[QUESTIONS_SHEET]
		else:
			questions_ws = wb.worksheets[0]

		question_rows = [list(row) for row in questions_ws.iter_rows(values_only=True)]

		settings_rows = []
		if SETTINGS_SHEET in wb.sheetnames:
			settings_rows = [list(row) for row in wb[SETTINGS_SHEET].iter_rows(values_only=True)]

		return question_rows, settings_rows
	finally:
		wb.close()


def add_error(errors, row, column, message):
	if len(errors) < MAX_ERRORS:
		errors.append({"row": row, "column": column, "message": message})


def read_question_header(rows, errors):
	"""Maps internal key -> column index. Unknown columns are ignored."""
	header = {}
	for index, value in enumerate(rows[0]):
		key = QUESTION_KEYS.get(header_key(value))
		if key and key not in header:
			header[key] = index

	missing = [label for label in REQUIRED_QUESTION_COLUMNS if QUESTION_KEYS[label.lower()] not in header]
	if missing:
		add_error(
			errors,
			1,
			None,
			_("Sheet '{0}' is missing these columns: {1}. Please use the downloaded template.").format(
				QUESTIONS_SHEET, ", ".join(missing)
			),
		)

	return header


def make_getter(values, header):
	def get(key):
		index = header.get(key)
		if index is None or index >= len(values):
			return ""
		return cell_value(values[index])

	return get


def read_settings_sheet(rows, errors):
	"""Quiz Title -> settings dict. An absent or empty sheet is fine."""
	if not rows:
		return {}

	header = {}
	for index, value in enumerate(rows[0]):
		fieldname = SETTINGS_FIELDS.get(header_key(value))
		if fieldname and fieldname not in header:
			header[fieldname] = index

	if "title" not in header:
		if any(any(cell_value(cell) for cell in row) for row in rows[1:]):
			add_error(
				errors,
				1,
				"Quiz Title",
				_("Sheet '{0}' needs a 'Quiz Title' column.").format(SETTINGS_SHEET),
			)
		return {}

	label_by_field = {fieldname: label for label, fieldname in SETTINGS_COLUMNS}
	settings = {}

	for index, values in enumerate(rows[1:]):
		if not any(cell_value(cell) for cell in values):
			continue

		row_number = index + 2
		get = make_getter(values, header)

		title = get("title")
		if not title:
			add_error(errors, row_number, "Quiz Title", _("Quiz Title is required."))
			continue

		parsed = {}
		for fieldname in header:
			if fieldname == "title":
				continue

			value = get(fieldname)
			if value == "":
				continue

			label = label_by_field[fieldname]
			if fieldname in SETTINGS_CHECKBOXES:
				parsed[fieldname] = to_bool(value)
			elif fieldname in SETTINGS_INTS:
				if not re.fullmatch(r"-?\d+", value):
					add_error(errors, row_number, label, _("{0} must be a whole number.").format(label))
					continue
				parsed[fieldname] = cint(value)
			else:
				parsed[fieldname] = value

		percentage = parsed.get("passing_percentage")
		if percentage is not None and (percentage < 0 or percentage > 100):
			add_error(
				errors,
				row_number,
				"Passing Percentage",
				_("Passing Percentage must be between 0 and 100."),
			)

		settings[title] = parsed

	return settings


def parse_row(values, header, row_number, errors):
	"""Row level and question type level validation. Returns a payload, or None."""
	get = make_getter(values, header)

	quiz_title = get("quiz_title")
	question = get("question")
	question_type = get("question_type")
	marks = get("marks")

	valid = True

	if not quiz_title:
		add_error(errors, row_number, "Quiz Title", _("Quiz Title is required."))
		valid = False

	if not plain_text(question):
		add_error(errors, row_number, "Question", _("Question is required."))
		valid = False

	matched_type = next((t for t in QUESTION_TYPES if t.lower() == question_type.lower()), None)
	if not matched_type:
		add_error(
			errors,
			row_number,
			"Question Type",
			_("Question Type must be one of: {0}.").format(", ".join(QUESTION_TYPES)),
		)
		valid = False

	if not marks:
		add_error(errors, row_number, "Marks", _("Marks is required."))
		valid = False
	elif not re.fullmatch(r"-?\d+", marks) or cint(marks) <= 0:
		add_error(errors, row_number, "Marks", _("Marks must be a whole number greater than 0."))
		valid = False

	if not valid:
		return None

	payload = {
		"row": row_number,
		"quiz_title": quiz_title,
		"question": question,
		"type": matched_type,
		"marks": cint(marks),
	}

	if matched_type == "Choices":
		valid = parse_choices(payload, get, row_number, errors)
	else:
		valid = reject_unused_columns(matched_type, get, row_number, errors)

		if matched_type == "User Input":
			valid = parse_expected_answer(payload, get, row_number, errors) and valid

	parse_explanation(payload, get, matched_type, row_number, errors)

	return payload if valid else None


def reject_unused_columns(question_type, get, row_number, errors):
	"""Filling a column that does not apply usually means the row is a mistake."""
	unused = [("Option {0}", f"option_{num}", num) for num in range(1, 5)]
	unused.append(("Correct Answer", "correct_answer", None))

	if question_type == "Open Ended":
		unused.append(("Expected Answer", "expected_answer", None))

	valid = True
	for label, key, num in unused:
		if get(key):
			column = label.format(num) if num else label
			add_error(
				errors,
				row_number,
				column,
				_("{0} does not apply to a '{1}' question. Please leave it empty.").format(
					column, question_type
				),
			)
			valid = False

	return valid


def parse_choices(payload, get, row_number, errors):
	valid = True
	options = []

	for num in range(1, 5):
		option = get(f"option_{num}")
		payload[f"option_{num}"] = option
		if option:
			options.append(option)

	if get("expected_answer"):
		add_error(
			errors,
			row_number,
			"Expected Answer",
			_("Expected Answer does not apply to a 'Choices' question. Please leave it empty."),
		)
		valid = False

	if not payload.get("option_1") or not payload.get("option_2"):
		add_error(
			errors,
			row_number,
			"Option 2",
			_("Minimum two options are required for multiple choice questions."),
		)
		valid = False

	if len(set(options)) != len(options):
		add_error(errors, row_number, "Option 1", _("Duplicate options found for this question."))
		valid = False

	return parse_correct_answer(payload, get, row_number, errors) and valid


def parse_correct_answer(payload, get, row_number, errors):
	raw = get("correct_answer")

	for num in range(1, 5):
		payload[f"is_correct_{num}"] = 0

	if not raw:
		add_error(
			errors,
			row_number,
			"Correct Answer",
			_("Correct Answer is required. Enter the correct option numbers, for example 1 or 1,3."),
		)
		return False

	valid = True
	seen = set()

	for part in re.split(r"[,;]", raw):
		part = part.strip()
		if not part:
			continue

		if not re.fullmatch(r"[1-4]", part):
			add_error(
				errors,
				row_number,
				"Correct Answer",
				_("Correct Answer must be option numbers between 1 and 4, for example 1 or 1,3."),
			)
			return False

		num = cint(part)
		if num in seen:
			continue
		seen.add(num)

		if not payload.get(f"option_{num}"):
			add_error(
				errors,
				row_number,
				"Correct Answer",
				_("Correct Answer refers to Option {0}, but Option {0} is empty.").format(num),
			)
			valid = False
			continue

		payload[f"is_correct_{num}"] = 1

	if not any(payload.get(f"is_correct_{num}") for num in range(1, 5)):
		add_error(
			errors,
			row_number,
			"Correct Answer",
			_("At least one option must be correct for this question."),
		)
		valid = False

	return valid


def parse_expected_answer(payload, get, row_number, errors):
	answers = split_multi(get("expected_answer"))

	for num in range(1, 5):
		payload[f"possibility_{num}"] = ""

	if not answers:
		add_error(
			errors,
			row_number,
			"Expected Answer",
			_("Expected Answer is required for a 'User Input' question."),
		)
		return False

	if len(answers) > 4:
		add_error(
			errors,
			row_number,
			"Expected Answer",
			_("At most 4 accepted answers are supported. Separate them with '{0}'.").format(MULTI_SEPARATOR),
		)
		return False

	for index, answer in enumerate(answers, start=1):
		payload[f"possibility_{index}"] = answer

	return True


def parse_explanation(payload, get, question_type, row_number, errors):
	"""'1: text | 3: text' for Choices, plain text otherwise."""
	for num in range(1, 5):
		payload.setdefault(f"explanation_{num}", "")

	raw = get("explanation")
	if not raw:
		return

	if question_type != "Choices":
		payload["explanation_1"] = raw
		return

	parts = split_multi(raw)
	reported = False

	for part in parts:
		match = re.match(r"^([1-4])\s*[:.)]\s*(.+)$", part, flags=re.DOTALL)
		if not match:
			# No option number given; a single unprefixed explanation belongs to option 1.
			if len(parts) == 1:
				payload["explanation_1"] = part
			elif not reported:
				reported = True
				add_error(
					errors,
					row_number,
					"Explanation",
					_("Write each explanation as '1: text' so it can be shown under the right option."),
				)
			continue

		num = cint(match.group(1))
		if not payload.get(f"option_{num}"):
			add_error(
				errors,
				row_number,
				"Explanation",
				_("Explanation refers to Option {0}, but Option {0} is empty.").format(num),
			)
			continue

		payload[f"explanation_{num}"] = match.group(2).strip()


def get_existing_quiz_state(title):
	"""Existing questions of a quiz, so the quiz level checks see the state after append."""
	quiz = frappe.db.get_value("LMS Quiz", {"title": title}, ["name", "limit_questions_to"], as_dict=True)
	if not quiz:
		return None

	rows = frappe.get_all(
		"LMS Quiz Question",
		filters={"parent": quiz.name, "parenttype": "LMS Quiz"},
		fields=["question", "marks"],
		order_by="idx asc",
	)

	texts = {}
	types = []
	if rows:
		names = [row.question for row in rows]
		for question in frappe.get_all(
			"LMS Question", filters={"name": ["in", names]}, fields=["name", "question", "type"]
		):
			texts[question.name] = question.question
			types.append(question.type)

	return {
		"name": quiz.name,
		"limit_questions_to": cint(quiz.limit_questions_to),
		"marks": [cint(row.marks) for row in rows],
		"types": types,
		"keys": {normalize(texts.get(row.question)) for row in rows if texts.get(row.question)},
	}


def validate_quiz_group(group, errors):
	"""Rules that only hold for the quiz as a whole, checked against the state after append."""
	existing = group["existing"]
	first_row = group["rows"][0]["row"] if group["rows"] else 1

	types = set(row["type"] for row in group["rows"])
	if existing:
		types |= set(existing["types"])

	if "Open Ended" in types and len(types) > 1:
		add_error(
			errors,
			first_row,
			"Question Type",
			_(
				"Quiz '{0}': if you want open ended questions then make sure each question in the quiz is of open ended type."
			).format(group["quiz_title"]),
		)

	limit = cint(existing["limit_questions_to"]) if existing else cint(group["settings"].get("limit_questions_to"))
	if not limit:
		return

	marks = [row["marks"] for row in group["rows"]]
	total = len(marks)
	if existing:
		marks += existing["marks"]
		total = len(existing["marks"]) + len(group["rows"])

	if limit >= total:
		add_error(
			errors,
			first_row,
			"Limit Questions To",
			_("Quiz '{0}': limit cannot be greater than or equal to the number of questions in the quiz.").format(
				group["quiz_title"]
			),
		)

	if len(set(marks)) > 1:
		add_error(
			errors,
			first_row,
			"Marks",
			_("Quiz '{0}': all questions should have the same marks if the limit is set.").format(
				group["quiz_title"]
			),
		)


def analyze(file_url):
	"""Runs every validation tier and returns (report, plan). Never writes."""
	question_rows, settings_rows = read_workbook(file_url)
	errors = []

	if not question_rows:
		frappe.throw(_("The file is empty."))

	header = read_question_header(question_rows, errors)
	if errors:
		return {"errors": errors, "quizzes": [], "skipped": [], "total_rows": 0}, []

	data_rows = [row for row in question_rows[1:] if any(cell_value(cell) for cell in row)]
	if not data_rows:
		frappe.throw(_("No question rows found in the file."))

	if len(data_rows) > MAX_ROWS:
		frappe.throw(_("The file has {0} rows. Maximum {1} rows are allowed.").format(len(data_rows), MAX_ROWS))

	quiz_settings = read_settings_sheet(settings_rows, errors)

	groups = {}
	skipped = []

	for index, values in enumerate(data_rows):
		row_number = index + 2
		payload = parse_row(values, header, row_number, errors)
		if not payload:
			continue

		title = payload["quiz_title"]
		group = groups.get(title)
		if not group:
			existing = get_existing_quiz_state(title)
			group = groups[title] = {
				"quiz_title": title,
				"existing": existing,
				# Settings only apply to a quiz being created.
				"settings": {} if existing else quiz_settings.get(title, {}),
				"rows": [],
				"keys": set(existing["keys"]) if existing else set(),
				"skipped": 0,
			}

		key = normalize(payload["question"])
		if key in group["keys"]:
			group["skipped"] += 1
			skipped.append(
				{"row": row_number, "quiz_title": title, "question": plain_text(payload["question"])[:120]}
			)
			continue

		group["keys"].add(key)
		group["rows"].append(payload)

	for group in groups.values():
		if group["rows"]:
			validate_quiz_group(group, errors)

	plan = [group for group in groups.values() if group["rows"]]

	report = {
		"errors": errors,
		"truncated": len(errors) >= MAX_ERRORS,
		"skipped": skipped,
		"total_rows": len(data_rows),
		"quizzes": [
			{
				"quiz_title": group["quiz_title"],
				"exists": bool(group["existing"]),
				"new_questions": len(group["rows"]),
				"skipped": group["skipped"],
			}
			for group in groups.values()
		],
	}

	return report, plan


@frappe.whitelist()
def validate_import(file_url):
	"""Dry run. Returns the report the modal shows before the user confirms."""
	check_permission()
	report, _plan = analyze(file_url)
	return report


def create_question(row):
	question = frappe.new_doc("LMS Question")
	question.question = row["question"]
	question.type = row["type"]

	if row["type"] == "Choices":
		for num in range(1, 5):
			question.set(f"option_{num}", row.get(f"option_{num}") or "")
			question.set(f"is_correct_{num}", cint(row.get(f"is_correct_{num}")))
	elif row["type"] == "User Input":
		for num in range(1, 5):
			question.set(f"possibility_{num}", row.get(f"possibility_{num}") or "")

	for num in range(1, 5):
		question.set(f"explanation_{num}", row.get(f"explanation_{num}") or "")

	question.insert()
	return question


def import_group(group):
	if group["existing"]:
		quiz = frappe.get_doc("LMS Quiz", group["existing"]["name"])
	else:
		quiz = frappe.new_doc("LMS Quiz")
		quiz.title = group["quiz_title"]
		for field, value in group["settings"].items():
			quiz.set(field, value)
		quiz.insert()

	for row in group["rows"]:
		question = create_question(row)
		quiz.append("questions", {"question": question.name, "marks": row["marks"]})

	quiz.save()
	return quiz


@frappe.whitelist()
def import_quiz(file_url):
	"""Validates again, then writes everything in a single transaction."""
	check_permission()

	report, plan = analyze(file_url)

	if report["errors"]:
		return {"success": False, **report}

	if not plan:
		return {"success": True, "imported": [], **report}

	imported = []
	for group in plan:
		quiz = import_group(group)
		imported.append(
			{
				"name": quiz.name,
				"quiz_title": quiz.title,
				"exists": bool(group["existing"]),
				"new_questions": len(group["rows"]),
				"skipped": group["skipped"],
			}
		)

	return {"success": True, "imported": imported, **report}
