import asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import uuid4

from docx import Document

from app.documents.repository import InMemoryDocumentRepository
from app.documents.schemas import DocumentKind as StoredDocumentKind
from app.documents.schemas import DocumentStatus, StoredDocument
from app.documents.source import CompositeDocumentSource
from app.documents.storage import InMemoryDocumentBlobStore
from app.extraction.schemas import DocumentKind


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Python and FastAPI")
    document.save(output)
    return output.getvalue()


def test_composite_source_resolves_clean_docx_and_rejects_quarantine() -> None:
    async def scenario() -> None:
        document_id = uuid4()
        documents = InMemoryDocumentRepository()
        blobs = InMemoryDocumentBlobStore()
        payload = _docx_bytes()
        stored = StoredDocument(
            id=document_id,
            owner_actor_id=uuid4(),
            organization_id=uuid4(),
            kind=StoredDocumentKind.CV,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            byte_size=len(payload),
            sha256="a" * 64,
            object_key="documents/clean",
            status=DocumentStatus.CLEAN,
            retention_until=datetime.now(UTC) + timedelta(days=1),
        )
        await documents.create(stored)
        await blobs.put(stored.object_key, payload, stored.content_type)
        source = CompositeDocumentSource(documents, blobs)

        resolved = await source.get(str(document_id), DocumentKind.CV)
        assert resolved.content == "Python and FastAPI"

        await documents.mark_status(document_id, DocumentStatus.QUARANTINED)
        try:
            await source.get(str(document_id), DocumentKind.CV)
        except ValueError as exc:
            assert str(exc) == "document_not_extractable"
        else:
            raise AssertionError("quarantined document was resolved")

    asyncio.run(scenario())
